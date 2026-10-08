from __future__ import annotations

import io
import hashlib
import json
import re
import threading
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .schema import Box, Item, PackingList, normalized_item_label, validate_for_export


ROOT = Path(__file__).resolve().parent.parent / ".local" / "training"
_lock = threading.Lock()
_SKU = re.compile(r"^([A-Z]{1,4}\s*[- ]?\s*\d{1,5})\s*(XXL|XL|XS|S|M|L)?\s*(.*)$", re.I)
_QTY = re.compile(r"\d+")


def _split_label(value: object) -> tuple[str, str, str]:
    label = str(value or "").strip()
    match = _SKU.match(label)
    if not match:
        return label, "", ""
    code = re.sub(r"^([A-Z]{1,4})\s*[- ]?\s*(\d+)$", r"\1-\2", match.group(1).upper())
    remainder = match.group(3).strip().strip("() ")
    return code, (match.group(2) or "").upper(), remainder


def parse_finished_workbook(raw: bytes, filename: str) -> PackingList:
    """Read the filled PRINT sheets, ignoring unused numbered template rows."""
    workbook = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
    sheets = [sheet for sheet in workbook if sheet.title.strip().upper().startswith("PRINT") or sheet.title.startswith("Page ")]
    if not sheets:
        raise ValueError("No PRINT or Page sheet found in this workbook")
    found: dict[int, Box] = {}
    private_mark = ""
    for sheet in sheets:
        headings: list[tuple[int, int, int]] = []
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().upper().replace(" ", "").startswith("BOXNO"):
                    number = sheet.cell(cell.row, cell.column + 1).value
                    if isinstance(number, (int, float)) and int(number) == number:
                        headings.append((cell.row, cell.column, int(number)))
        for heading_row, column, number in headings:
            next_rows = [row for row, col, _ in headings if col == column and row > heading_row]
            end_row = min(next_rows) if next_rows else sheet.max_row + 1
            header = None
            for row in range(heading_row + 1, min(heading_row + 9, end_row)):
                value = str(sheet.cell(row, column).value or "").strip().upper()
                if value == "ITEM CODE":
                    header = row
                    break
            if header is None:
                continue
            mark = sheet.cell(heading_row + 2, column + 1).value
            if mark and not private_mark:
                private_mark = str(mark).strip()
            items: list[Item] = []
            for row in range(header + 1, end_row):
                label = sheet.cell(row, column).value
                quantity = sheet.cell(row, column + 1).value
                if not label or quantity is None:
                    continue
                number_match = _QTY.search(str(quantity))
                if not number_match:
                    continue
                code, size, note = _split_label(label)
                items.append(Item(code=code, size=size, note=note, quantity=int(number_match.group())))
            if items:
                if number in found:
                    raise ValueError(f"Box {number} occurs more than once in the workbook")
                found[number] = Box(number=number, items=items)
    if not found:
        raise ValueError("No filled boxes were found in the PRINT sheet")
    customer = re.split(r"\b\d{2}[- ]\d{2}[- ]\d{4}\b", Path(filename).stem, maxsplit=1)[0].strip(" -_")
    date_match = re.search(r"\b\d{2}[- ]\d{2}[- ]\d{4}\b", Path(filename).stem)
    result = PackingList(customer=customer, packing_date=date_match.group().replace(" ", "-") if date_match else "", private_mark=private_mark, boxes=sorted(found.values(), key=lambda box: box.number))
    result.warnings.append("Imported from the finished Excel. Compare entries with the photos, then press Final to save it as a verified example.")
    return result


def save_verified_case(data: PackingList, photos: list[bytes], origin: str) -> str:
    validate_for_export(data)
    if not photos:
        raise ValueError("At least one source photo is required to save a verified example")
    case_id = uuid.uuid4().hex
    with _lock:
        folder = ROOT / case_id
        folder.mkdir(parents=True)
        record = {"id": case_id, "created_at": datetime.now(timezone.utc).isoformat(), "origin": origin,
                  "list": data.model_dump(), "photos": [f"page-{i + 1}.jpg" for i in range(len(photos))],
                  "photo_hashes": [hashlib.sha256(photo).hexdigest() for photo in photos]}
        for index, photo in enumerate(photos, start=1):
            (folder / f"page-{index}.jpg").write_bytes(photo)
        (folder / "case.json").write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")
    return case_id


def replace_verified_case(case_id: str, data: PackingList) -> None:
    validate_for_export(data)
    path = ROOT / case_id / "case.json"
    with _lock:
        record = json.loads(path.read_text(encoding="utf-8"))
        record["list"] = data.model_dump()
        record["updated_at"] = datetime.now(timezone.utc).isoformat()
        path.write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")


def _records() -> list[dict]:
    if not ROOT.exists():
        return []
    records = []
    for path in ROOT.glob("*/case.json"):
        try:
            records.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            continue
    return sorted(records, key=lambda record: record.get("created_at", ""), reverse=True)


def find_exact_case(photos: list[bytes]) -> tuple[str, PackingList] | None:
    hashes = [hashlib.sha256(photo).hexdigest() for photo in photos]
    for record in _records():
        previous = record.get("photo_hashes")
        if previous is None:
            folder = ROOT / record["id"]
            previous = [hashlib.sha256((folder / name).read_bytes()).hexdigest() for name in record.get("photos", [])]
        if previous == hashes:
            return record["id"], PackingList.model_validate(record["list"])
    return None


def training_stats() -> dict:
    records = _records()
    codes = {item.get("code", "").upper() for record in records for box in record["list"]["boxes"] for item in box["items"]}
    return {"verified_examples": len(records), "known_codes": len(codes), "verified_boxes": sum(len(record["list"]["boxes"]) for record in records)}


def verified_guidance() -> str:
    """Keep prompt hints compact so saved labels guide recognition without swamping the image."""
    records = _records()
    if not records:
        return ""
    codes: Counter[str] = Counter()
    labels: Counter[str] = Counter()
    for record in records:
        for box in record["list"]["boxes"]:
            for raw_item in box["items"]:
                item = Item.model_validate(raw_item)
                codes[item.code.upper()] += 1
                labels[normalized_item_label(item)] += 1
    code_hints = ", ".join(code for code, _ in codes.most_common(85) if code)
    example_labels = ", ".join(label for label, _ in labels.most_common(35) if label)
    return ("\nVERIFIED COMPANY EXAMPLES (use only as spelling/code hints; the current photo is authoritative):\n"
            f"Known product codes/names: {code_hints[:1300]}\n"
            f"Previously confirmed item labels: {example_labels[:1200]}\n"
            "Never insert a known item merely because it appears here. Read the image and flag ambiguous writing.\n")
