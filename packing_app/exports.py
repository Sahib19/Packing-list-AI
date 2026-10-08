from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from .schema import Box, PackingList, normalized_item_label, validate_for_export


MIN_ROWS_PER_BOX = 5
ROW_PT = 22
GROUP_PT = 74
GROUP_GAP_PT = 16
A4_GROUP_BUDGET_PT = 440


def visible_rows(box: Box) -> int:
    return max(MIN_ROWS_PER_BOX, len(box.items))


def group_height(group: list[Box]) -> int:
    return GROUP_PT + ROW_PT * max(visible_rows(box) for box in group) + GROUP_GAP_PT


def plan_pages(boxes: list[Box]) -> list[list[list[Box]]]:
    """Pair boxes side by side, then pack whole pairs onto pages."""
    ordered = sorted(boxes, key=lambda box: box.number)
    pairs = [ordered[i:i + 2] for i in range(0, len(ordered), 2)]
    pages: list[list[list[Box]]] = []
    current: list[list[Box]] = []
    height = 0
    for pair in pairs:
        required = group_height(pair)
        if current and (len(current) == 2 or height + required > A4_GROUP_BUDGET_PT):
            pages.append(current)
            current = []
            height = 0
        current.append(pair)
        height += required
    if current:
        pages.append(current)
    return pages


def _page_size(groups: list[list[Box]]) -> tuple[float, float]:
    needed = sum(group_height(group) for group in groups)
    if needed <= A4_GROUP_BUDGET_PT:
        return landscape(A4)
    a3_width, a3_height = landscape(A3)
    if needed <= a3_height - 120:
        return a3_width, a3_height
    return a3_width, needed + 120


def _sheet_box(ws, box: Box | None, left_col: int, top_row: int, rows: int, private_mark: str) -> int:
    last_col = left_col + 2
    if box is not None:
        ws.cell(top_row, left_col + 1, "BOX No.")
        ws.cell(top_row, last_col, box.number)
        ws.cell(top_row + 1, left_col + 1, "PVT. MARK")
        ws.cell(top_row + 1, last_col, private_mark)
    labels = ("S. No.", "ITEM CODE", "QTY.")
    for offset, label in enumerate(labels):
        cell = ws.cell(top_row + 2, left_col + offset, label)
        cell.fill = PatternFill("solid", fgColor="E9EDF2")
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    edge = Side(style="thin", color="AAB4C0")
    for index in range(rows):
        row = top_row + 3 + index
        item = box.items[index] if box is not None and index < len(box.items) else None
        if box is not None:
            ws.cell(row, left_col, index + 1)
        if item is not None:
            ws.cell(row, left_col + 1, normalized_item_label(item))
            ws.cell(row, last_col, item.quantity)
            ws.cell(row, last_col).number_format = '0" PCS"'
        for col in range(left_col, last_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", horizontal="center" if col != left_col + 1 else "left")
            cell.border = Border(bottom=edge)
        ws.row_dimensions[row].height = ROW_PT
    for row in (top_row, top_row + 1, top_row + 2):
        ws.row_dimensions[row].height = 23
    for col in (left_col, last_col):
        ws.cell(top_row, col).font = Font(name="Arial", size=11, bold=True)
    return top_row + 3 + rows


def make_excel(data: PackingList) -> bytes:
    validate_for_export(data)
    wb = Workbook()
    wb.remove(wb.active)
    pages = plan_pages(data.boxes)
    for page_index, groups in enumerate(pages, start=1):
        ws = wb.create_sheet(f"Page {page_index:02d}")
        ws.sheet_view.showGridLines = False
        for col, width in {"A": 7, "B": 34, "C": 11, "D": 3, "E": 7, "F": 34, "G": 11}.items():
            ws.column_dimensions[col].width = width
        ws.merge_cells("A1:G1")
        title = ws["A1"]
        title.value = "PACKING LIST"
        title.font = Font(name="Arial", size=16, bold=True)
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        ws["A2"] = "TO:"
        ws["B2"] = data.customer
        ws["E2"] = "DATE:"
        ws["F2"] = data.packing_date
        ws["A3"] = "TOTAL BOXES:"
        ws["B3"] = len(data.boxes)
        ws["E3"] = "TRANSPORT:"
        ws["F3"] = data.transport
        for row in (2, 3):
            ws.row_dimensions[row].height = 23
            for col in range(1, 8):
                ws.cell(row, col).font = Font(name="Arial", size=10, bold=col in (1, 5))
        next_row = 5
        for group in groups:
            rows = max(visible_rows(box) for box in group)
            left_end = _sheet_box(ws, group[0], 1, next_row, rows, data.private_mark)
            _sheet_box(ws, group[1] if len(group) > 1 else None, 5, next_row, rows, data.private_mark)
            next_row = left_end + 2
        ws.print_area = f"A1:G{next_row - 2}"
        ws.print_options.horizontalCentered = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if sum(group_height(group) for group in groups) > A4_GROUP_BUDGET_PT else ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0, footer=0)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _fit_text(c: Canvas, text: str, x: float, y: float, width: float, font: str = "Helvetica", size: float = 9.2) -> None:
    text = text or ""
    while size > 6.5 and stringWidth(text, font, size) > width:
        size -= 0.3
    if stringWidth(text, font, size) > width:
        while text and stringWidth(text + "…", font, size) > width:
            text = text[:-1]
        text += "…"
    c.setFont(font, size)
    c.drawString(x, y, text)


def _pdf_box(c: Canvas, box: Box | None, x: float, top: float, width: float, rows: int, private_mark: str) -> None:
    c.setStrokeColor(colors.HexColor("#8794A4"))
    c.setLineWidth(0.6)
    if box is None:
        return
    c.setFont("Helvetica-Bold", 10)
    c.drawString(x + 7, top - 14, "BOX No.")
    c.drawRightString(x + width - 7, top - 14, str(box.number))
    c.setFont("Helvetica", 9)
    c.drawString(x + 7, top - 35, "PVT. MARK")
    _fit_text(c, private_mark, x + 90, top - 35, width - 97, size=9)
    header_top = top - 48
    header_bottom = header_top - 21
    c.setFillColor(colors.HexColor("#E9EDF2"))
    c.rect(x, header_bottom, width, 21, fill=1, stroke=0)
    c.setFillColor(colors.black)
    number_w, qty_w = 43, 76
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(x + number_w / 2, header_bottom + 6, "S. No.")
    c.drawString(x + number_w + 8, header_bottom + 6, "ITEM CODE")
    c.drawCentredString(x + width - qty_w / 2, header_bottom + 6, "QTY.")
    table_bottom = header_bottom - ROW_PT * rows
    for xpos in (x, x + number_w, x + width - qty_w, x + width):
        c.line(xpos, header_top, xpos, table_bottom)
    for line_index in range(rows + 1):
        y = header_bottom - line_index * ROW_PT
        c.line(x, y, x + width, y)
    c.line(x, header_top, x + width, header_top)
    for index in range(rows):
        y = header_bottom - (index + 1) * ROW_PT + 7
        c.setFont("Helvetica", 9)
        c.drawCentredString(x + number_w / 2, y, str(index + 1))
        if index < len(box.items):
            item = box.items[index]
            _fit_text(c, normalized_item_label(item), x + number_w + 7, y, width - number_w - qty_w - 14)
            if item.quantity is not None:
                c.drawCentredString(x + width - qty_w / 2, y, f"{item.quantity} PCS")


def make_pdf(data: PackingList) -> bytes:
    validate_for_export(data)
    output = BytesIO()
    pages = plan_pages(data.boxes)
    canvas = Canvas(output, pagesize=landscape(A4))
    for page_index, groups in enumerate(pages, start=1):
        page_width, page_height = _page_size(groups)
        canvas.setPageSize((page_width, page_height))
        margin = 26
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawCentredString(page_width / 2, page_height - 30, "PACKING LIST")
        canvas.setFont("Helvetica", 9.5)
        _fit_text(canvas, f"TO: {data.customer}", margin, page_height - 51, page_width * 0.55, size=9.5)
        canvas.drawRightString(page_width - margin, page_height - 51, f"DATE: {data.packing_date}")
        canvas.drawString(margin, page_height - 67, f"TOTAL BOXES: {len(data.boxes)}")
        canvas.drawRightString(page_width - margin, page_height - 67, f"TRANSPORT: {data.transport}")
        canvas.setStrokeColor(colors.HexColor("#B8C1CC"))
        canvas.line(margin, page_height - 76, page_width - margin, page_height - 76)
        gap = 15
        box_width = (page_width - 2 * margin - gap) / 2
        top = page_height - 89
        for group in groups:
            rows = max(visible_rows(box) for box in group)
            _pdf_box(canvas, group[0], margin, top, box_width, rows, data.private_mark)
            _pdf_box(canvas, group[1] if len(group) > 1 else None, margin + box_width + gap, top, box_width, rows, data.private_mark)
            top -= group_height(group)
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(page_width - margin, 14, f"Page {page_index} of {len(pages)}")
        canvas.showPage()
    canvas.save()
    return output.getvalue()

